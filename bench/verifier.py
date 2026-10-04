#!/usr/bin/env python3
"""Le vérificateur de Visa : relit le livrable d'un agent contre les pièces, sans jamais voir la grille.

  python3 bench/verifier.py harvey-labs/results/<tâche>/<source>-<condition>/<horodatage>
  python3 bench/verifier.py <run_id> --modele mistral-medium-3.5

Il écrit revue-visa.md (pour l'agent), revue-visa.json (pour la boucle) et revue-legora.txt (à coller)
dans le dossier du lancement. Trois recherches :
- ce qui manque : un relecteur Mistral établit, depuis la pièce de référence, la liste des éléments à
  traiter (comme un tableau de concordance), puis dit lesquels le livrable ne traite pas ;
- ce qui est faux : chaque chiffre, date et section du livrable est cherché dans les pièces par script ;
  les introuvables et les affirmations contredites passent au relecteur ;
- les calculs : totaux des tableurs et opérations écrites (« A − B = C ») refaits par script, et les
  calculs que le relecteur repère, recalculés par script.
Chaque problème porte un extrait des pièces retrouvé mot pour mot par script ; sinon il est écarté.
Le vérificateur ne lit que la consigne, les pièces et le livrable.
"""
import argparse
import http.client
import json
import math
import re
import sys
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import run as banc  # noqa: E402  (load_env, ROOT, LAB)

SCRIPTS = banc.ROOT / "skills" / "cross-document-review" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from extract_text import convert  # noqa: E402
from sections import blocks, locate, normalize, unit_numbers  # noqa: E402

MODELE = "mistral-large-latest"
# Tarifs publics Mistral (USD par million de jetons, entrée / sortie), pour une estimation du coût.
TARIFS = {"mistral-large": (0.5, 1.5), "mistral-medium": (0.4, 2.0)}
MOTS_PAR_MORCEAU = 3500
CONSIGNE_A_TELEVERSER = Path.home() / "lab-claude-code" / "a-televerser"
JETONS = {"entree": 0, "sortie": 0, "appels": 0}
ECARTES: list[dict] = []  # ce que le relecteur a proposé et que le script a rejeté, pour comprendre


# ---------------------------------------------------------------- Mistral

def mistral(systeme: str, message: str, modele: str, max_tokens: int = 16000) -> dict:
    cle = banc.load_env().get("MISTRAL_API_KEY")
    if not cle:
        sys.exit("MISTRAL_API_KEY manque dans .env")
    corps = json.dumps({"model": modele, "temperature": 0, "top_p": 1, "max_tokens": max_tokens,
                        "response_format": {"type": "json_object"},
                        "messages": [{"role": "system", "content": systeme}, {"role": "user", "content": message}]}).encode()
    attente = 5.0
    for essai in range(7):
        requete = urllib.request.Request("https://api.mistral.ai/v1/chat/completions", data=corps,
                                         headers={"Authorization": f"Bearer {cle}", "Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(requete, timeout=900) as r:
                reponse = json.loads(r.read())
            break
        except urllib.error.HTTPError as e:
            if e.code not in (429, 500, 502, 503, 504) or essai == 6:
                raise RuntimeError(f"Mistral {e.code} : {e.read()[:300]!r}") from None
        except (urllib.error.URLError, OSError, http.client.HTTPException) as e:  # coupure, délai, connexion réinitialisée
            if essai == 6:
                raise RuntimeError(f"Mistral injoignable : {e}") from None
        time.sleep(min(attente, 90))
        attente *= 2
    usage = reponse.get("usage", {})
    JETONS["entree"] += usage.get("prompt_tokens", 0)
    JETONS["sortie"] += usage.get("completion_tokens", 0)
    JETONS["appels"] += 1
    texte = reponse["choices"][0]["message"]["content"]
    try:
        return json.loads(texte)
    except json.JSONDecodeError:
        debut, fin = texte.find("{"), texte.rfind("}")
        try:
            return json.loads(texte[debut:fin + 1])
        except json.JSONDecodeError:
            return {}


# ---------------------------------------------------------------- pièces et livrable

def texte_docx_cellules(chemin: Path) -> str:
    """Paragraphes et cellules de tableau un par ligne : les citations prises dans un tableau s'y retrouvent."""
    import docx

    d = docx.Document(str(chemin))
    lignes = [p.text for p in d.paragraphs]
    for table in d.tables:
        for rangee in table.rows:
            vues = set()
            for cellule in rangee.cells:
                if id(cellule._tc) not in vues:
                    vues.add(id(cellule._tc))
                    lignes.append(cellule.text)
    return "\n".join(lignes)


def lire(chemin: Path) -> dict:
    texte = convert(chemin)
    variantes = [texte]
    if chemin.suffix.lower() == ".docx":
        try:
            variantes.append(texte_docx_cellules(chemin))
        except Exception:
            pass
    return {"texte": texte, "normes": [normalize(v) for v in variantes]}


GUILLEMETS = ' "“”«»\'‘’'


def fragments(citation: str) -> list[str]:
    """Une citation sans ses guillemets, coupée à ses « ... » : chaque morceau doit être mot pour mot."""
    brut = " ".join(str(citation or "").split()).strip(GUILLEMETS)
    morceaux = [m.strip(GUILLEMETS + ".,;:") for m in re.split(r"\[?(?:\.\.\.|…)\]?", brut)]
    return [m for m in morceaux if len(normalize(m).split()) >= 2]


def retrouve(citation: str, doc: dict) -> bool:
    morceaux = fragments(citation)
    if not morceaux or max(len(normalize(m).split()) for m in morceaux) < 3:
        return False
    return any(all(locate(m, n)[0] == "exact" for m in morceaux) for n in doc["normes"])


def ou_est(citation: str, docs: dict, prefere: str = "") -> str | None:
    """Le nom de la pièce où la citation se trouve mot pour mot (la pièce annoncée d'abord), ou None."""
    noms = sorted(docs, key=lambda n: n != prefere)
    return next((n for n in noms if retrouve(citation, docs[n])), None)


def nom_proche(nom: str, docs: dict) -> str:
    nom = (nom or "").strip().split("/")[-1]
    if nom in docs:
        return nom
    return next((n for n in docs if Path(n).stem.lower() == Path(nom).stem.lower()), nom)


# ---------------------------------------------------------------- calculs refaits par script

def nombre(brut) -> float | None:
    if isinstance(brut, (int, float)) and not isinstance(brut, bool):
        return float(brut)
    s = str(brut or "").strip()
    negatif = s.startswith("(") and s.endswith(")") or s.startswith(("-", "−"))
    s = re.sub(r"[^\d.]", "", s.replace(",", ""))
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if negatif else v


MONTANT = re.compile(r"^\(?[-−]?[$€£]?\s?\d[\d,]*(?:\.\d+)?\)?$")


def num(cellule) -> float | None:
    """Une cellule de tableur qui porte un nombre, même écrit en texte (« 4,800,000 », « (400,000) »)."""
    if isinstance(cellule, bool):
        return None
    if isinstance(cellule, (int, float)):
        return float(cellule)
    if isinstance(cellule, str) and MONTANT.match(cellule.strip()):
        return nombre(cellule)
    return None


def libelle_de(rangee) -> str:
    return next((str(c).strip() for c in rangee if isinstance(c, str) and c.strip() and num(c) is None), "")


def verifier_bloc(lignes: list[tuple[str, list]], nom: str, ou: str) -> list[dict]:
    """lignes : (libellé, valeurs par colonne). Une ligne « Total » doit valoir la somme des lignes simples
    du bloc au-dessus (ou celle des sous-totaux). Prudent : on ne signale que si aucune des deux ne tombe juste."""
    ecarts = []
    for i, (libelle, valeurs) in enumerate(lignes):
        if not libelle or not re.search(r"\btotal\b", libelle, re.I):
            continue
        for j, annonce in enumerate(valeurs):
            if annonce is None:
                continue
            simples, sous_totaux, operandes = [], [], []
            for lib, vals in reversed(lignes[:i]):
                if lib is None:  # intertitre ou ligne vide qui ferme le bloc
                    if simples or sous_totaux:
                        break
                    continue
                v = vals[j] if j < len(vals) else None
                if v is None:
                    continue
                if re.search(r"\btotal\b", lib, re.I):
                    sous_totaux.append(v)
                else:
                    simples.append(v)
                    operandes.append((lib, v))
            if len(simples) < 2:
                continue
            # Les k lignes les plus proches (un bloc qui commence plus bas), ou des sous-totaux de la feuille
            # (« TOTAL ASSETS » = actif courant + actif non courant).
            avant = [vals[j] for lib, vals in reversed(lignes[:i]) if lib and re.search(r"\btotal\b", lib, re.I)
                     and j < len(vals) and vals[j] is not None]
            justes = ([sum(simples[:k]) for k in range(2, len(simples) + 1)] + [sum(avant[:k]) for k in range(2, len(avant) + 1)]
                      + [sum(sous_totaux), sum(simples) + sum(sous_totaux)])
            if not any(math.isclose(annonce, s, abs_tol=1.0) for s in justes if s):
                ecarts.append({"piece": nom, "ou": ou, "libelle": libelle, "annonce": annonce, "recalcule": sum(simples),
                               "operandes": list(reversed(operandes))})
    return ecarts


def totaux_tableurs(chemin: Path) -> list[dict]:
    import openpyxl

    ecarts = []
    for feuille in openpyxl.load_workbook(chemin, data_only=True).worksheets:
        lignes = []
        for rangee in feuille.iter_rows(values_only=True):
            valeurs = [num(c) for c in rangee]
            libelle = libelle_de(rangee)
            lignes.append((libelle, valeurs) if any(v is not None for v in valeurs) else (None, []))
        ecarts += verifier_bloc(lignes, chemin.name, f"feuille « {feuille.title} »")
    return ecarts


CHIFFRE_LIGNE = re.compile(r"\(?[-−]?[$€£]\s?\d[\d,]*(?:\.\d+)?\)?|\(?\b\d{1,3}(?:,\d{3})+(?:\.\d+)?\)?")


def totaux_texte(nom: str, texte: str) -> list[dict]:
    """Les tableaux d'un texte (docx converti) : « libellé  $montant » ligne à ligne, puis « Total … $montant »."""
    lignes = []
    for ligne in texte.splitlines():
        brut = ligne.strip().strip("|").strip()
        if not brut or re.fullmatch(r"[-=+|: ]+", brut):
            continue
        montants = [nombre(m.group(0)) for m in CHIFFRE_LIGNE.finditer(brut)]
        libelle = CHIFFRE_LIGNE.sub("", brut).strip(" |:")
        if montants and len(libelle.split()) <= 14:
            lignes.append((libelle, montants))
        else:
            lignes.append((None, []))
    return verifier_bloc(lignes, nom, "tableau")


OPERATION = re.compile(r"(\(?[$€£]?\s?\d[\d,]*(?:\.\d+)?\)?(?:\s*(?:[−\-–+×x*/])\s*\(?[$€£]?\s?\d[\d,]*(?:\.\d+)?\)?)+)"
                       r"\s*=\s*\(?[$€£]?\s?(\d[\d,]*(?:\.\d+)?)\)?")


def evaluer(expression: str) -> float | None:
    propre = (expression.replace("−", "-").replace("–", "-").replace("×", "*").replace(",", "")
              .replace("$", "").replace("€", "").replace("£", ""))
    propre = re.sub(r"(?<=\d)\s*x\s*(?=[\d(])", "*", propre)
    if not re.fullmatch(r"[\d.\s+\-*/()%]+", propre):
        return None
    propre = re.sub(r"(\d+(?:\.\d+)?)\s*%", r"(\1/100)", propre)
    try:
        return float(eval(propre, {"__builtins__": {}}, {}))  # noqa: S307 (chiffres et opérateurs seulement)
    except Exception:
        return None


def operations_ecrites(nom: str, texte: str) -> list[dict]:
    ecarts = []
    for m in OPERATION.finditer(texte):
        calcul, annonce = evaluer(m.group(1)), nombre(m.group(2))
        if calcul is None or annonce is None:
            continue
        if not math.isclose(abs(calcul), annonce, rel_tol=0.005, abs_tol=1.0):
            ecarts.append({"piece": nom, "extrait": " ".join(m.group(0).split()), "recalcule": calcul, "annonce": annonce})
    return ecarts


def plausible(annonce: float, calcul: float) -> bool:
    """Un vrai écart de calcul est une erreur de quelques pour cent (4,8 M au lieu de 4,9 M). Un résultat de signe
    opposé ou loin du montant annoncé signale plutôt que la formule proposée par le relecteur est fausse."""
    return (annonce >= 0) == (calcul >= 0) and abs(calcul - annonce) <= 0.25 * max(abs(annonce), 1.0)


def arrondi_compatible(annonce: float, calcul: float) -> bool:
    """« $2.6 million » pour 2 612 500 : juste à la précision affichée."""
    if math.isclose(abs(annonce), abs(calcul), rel_tol=0.002, abs_tol=1.0):
        return True
    for chiffres in (1, 2, 3):
        if abs(calcul) >= 1e5 and math.isclose(abs(annonce), round(abs(calcul), -int(math.log10(abs(calcul))) + chiffres),
                                                rel_tol=1e-9):
            return True
    return False


# ---------------------------------------------------------------- ce qui manque

SYSTEME = ("Tu es un avocat d'affaires senior qui relit le travail d'un collaborateur, de façon indépendante. "
           "Tu réponds uniquement en JSON valide. Toute citation est copiée MOT POUR MOT depuis le texte fourni "
           "(5 à 40 mots, sans « ... », sans reformuler, sans traduire) : une citation inexacte sera rejetée par un script.")


def roles(consigne: str, docs: dict, modele: str) -> dict:
    apercu = "\n\n".join(f"### {n} ({len(d['texte'].split())} mots)\n{' '.join(d['texte'].split()[:250])}"
                         for n, d in docs.items())
    r = mistral(SYSTEME, f"""Consigne donnée à l'agent :
<consigne>{consigne}</consigne>

Pièces (début de chacune) :
{apercu}

Quelle(s) pièce(s) font foi selon la consigne (la référence contre laquelle on vérifie : en général l'acte
juridique qui gouverne — contrat, accord signé, loi —, pas la liste, le tableau de suivi ou l'annexe qu'on vérifie) ? Lesquelles sont
vérifiées (les sujets) ? Lesquelles servent de contexte ? Qui est le client et de quel côté ?
Réponds : {{"reference": ["nom exact"], "sujets": ["nom exact"], "contexte": ["nom exact"], "client": "…"}}""", modele, 2000)
    r["reference"] = [nom_proche(n, docs) for n in r.get("reference") or [] if nom_proche(n, docs) in docs] or [max(
        docs, key=lambda n: len(docs[n]["texte"].split()))]
    r["sujets"] = [nom_proche(n, docs) for n in r.get("sujets") or [] if nom_proche(n, docs) in docs
                   and nom_proche(n, docs) not in r["reference"]] or [n for n in docs if n not in r["reference"]]
    return r


def morceaux(texte: str) -> list[str]:
    sortie, courant, compte = [], [], 0
    for _, bloc in blocks(texte):
        courant.append(bloc)
        compte += len(bloc.split())
        if compte >= MOTS_PAR_MORCEAU:
            sortie.append("\n\n".join(courant))
            courant, compte = [], 0
    if courant:
        sortie.append("\n\n".join(courant))
    return sortie


def elements_du_morceau(consigne: str, ref: str, morceau: str, n: int, total: int, autres: dict, role: dict,
                        modele: str) -> list[dict]:
    pieces = "\n\n".join(f'<piece nom="{nom}">\n{d["texte"]}\n</piece>' for nom, d in autres.items())
    r = mistral(SYSTEME, f"""Consigne donnée à l'agent (tu ne vois pas son livrable) :
<consigne>{consigne}</consigne>
Client : {role.get('client', 'non précisé')}

Les autres pièces du dossier :
{pieces}

Voici la partie {n}/{total} de la pièce de référence « {ref} » :
<reference nom="{ref}">
{morceau}
</reference>

Fais le tableau de concordance de CETTE partie, ligne par ligne : chaque exigence, obligation, livrable,
condition, déclaration, définition chiffrée, délai, montant ou seuil que la consigne demande de vérifier
devient un élément. Pour chacun, cherche ce que disent les autres pièces et classe-le :
- "conforme" : les autres pièces le reprennent fidèlement ;
- "ecart" : une autre pièce le reprend avec une valeur, une partie, un délai, un seuil ou une portée différente ;
- "absent" : aucune autre pièce ne le traite alors que, vu la consigne, l'une d'elles devrait le reprendre ;
- "incoherence" : deux pièces se contredisent entre elles à son sujet ;
- "hors_champ" : les autres pièces n'ont pas vocation à le reprendre (par exemple une déclaration générale qu'une
  liste de closing ou un tableau de suivi ne reprend jamais) — ne le liste pas.
Sois exhaustif sur cette partie (n'oublie ni les définitions, ni les annexes, ni les tableaux, ni les notes).
Ignore le pur boilerplate (notifications, signatures, interprétation).
Donne à chaque élément une importance : "haute" (montant, pourcentage, délai, date, partie, condition, consentement,
livrable de closing, risque juridique ou financier réel pour le client), "moyenne" (portée, qualification, renvoi de
section), "basse" (rédaction, détail sans conséquence).

Réponds : {{"elements": [{{"section": "Section 4.06(b)", "sujet": "titre court en français",
"citation_reference": "copie mot pour mot de la référence", "constat": "conforme|ecart|absent|incoherence",
"importance": "haute|moyenne|basse",
"piece_comparee": "nom exact ou vide", "citation_piece": "copie mot pour mot de la pièce comparée, ou vide si absent",
"detail": "en français : les deux valeurs et la différence, ou pourquoi c'est conforme"}}]}}""", modele)
    return r.get("elements") or []


def traitement(consigne: str, elements: list[dict], livrable: str, modele: str) -> dict:
    liste = "\n".join(f"[{e['id']}] {e.get('section', '')} — {e.get('sujet', '')} — constat : {e.get('constat')} — "
                      f"{e.get('detail', '')}" for e in elements)
    r = mistral(SYSTEME, f"""Consigne donnée à l'agent :
<consigne>{consigne}</consigne>

Livrable de l'agent :
<livrable>
{livrable}
</livrable>

Voici des éléments qu'un relecteur indépendant a relevés dans les pièces. Pour CHACUN, dis si le livrable
le traite : "oui" (il le traite correctement, écart signalé avec les bonnes valeurs le cas échéant),
"partiel" (il en parle mais il manque l'écart, une valeur, la conséquence ou la recommandation), "non" (il
n'en parle pas). Un élément conforme compte comme traité s'il est mentionné ou couvert par la section du
livrable qui lui correspond. Cherche bien dans tout le livrable, tableaux et annexes compris, avant de dire "non".

{liste}

Réponds : {{"verdicts": [{{"id": "E-001", "traite": "oui|partiel|non",
"citation_livrable": "copie mot pour mot du livrable si oui ou partiel", "manque": "en français : ce qu'il faut ajouter ou corriger"}}]}}""",
                modele)
    return {v.get("id"): v for v in r.get("verdicts") or []}


def dans_le_champ(section: str, sujets_norm: str) -> bool:
    """Un élément absent n'est bloquant que si la pièce vérifiée traite de sa section (« 8.1 » pour « 8.1(k) ») :
    une déclaration générale (fiscalité, environnement) qu'une liste de closing ne reprend jamais ne compte pas."""
    m = re.search(r"\d+\.\d+", section or "")
    return not m or re.search(rf"(?<![\d.]){re.escape(m.group(0))}(?![\d])", sujets_norm) is not None


def ce_qui_manque(consigne: str, docs: dict, livrable: str, role: dict, modele: str) -> tuple[list[dict], dict]:
    taches = []
    parcourues = role["reference"] + [n for n in role["sujets"] if n not in role["reference"]]
    for ref in parcourues:
        autres = {n: d for n, d in docs.items() if n != ref}
        parts = morceaux(docs[ref]["texte"])
        taches += [(ref, m, i + 1, len(parts), autres) for i, m in enumerate(parts)]
    with ThreadPoolExecutor(max_workers=6) as pool:
        resultats = list(pool.map(lambda t: (t[0], elements_du_morceau(consigne, t[0], t[1], t[2], t[3], t[4], role, modele)),
                                  taches))
    elements, ecartes = [], 0
    for ref, liste in resultats:
        for e in liste:
            if e.get("constat") not in ("conforme", "ecart", "absent", "incoherence"):
                continue
            if not ou_est(e.get("citation_reference", ""), {ref: docs[ref]}):
                ecartes += 1
                ECARTES.append({"raison": "citation de la référence introuvable", "ref": ref, **e})
                continue
            e["reference"] = ref
            citee = (e.get("citation_piece") or "").strip()
            if citee:
                trouve = ou_est(citee, docs, nom_proche(e.get("piece_comparee", ""), docs))
                if trouve:
                    e["piece_comparee"] = trouve
                else:
                    e["citation_piece"] = ""
                    if e.get("constat") in ("ecart", "incoherence"):
                        ecartes += 1  # un écart sans extrait vérifié de l'autre pièce ne tient pas
                        ECARTES.append({"raison": "citation de la pièce comparée introuvable", **e})
                        continue
            e["id"] = f"E-{len(elements) + 1:03d}"
            elements.append(e)
    lots = [elements[i:i + 40] for i in range(0, len(elements), 40)]
    verdicts: dict = {}
    with ThreadPoolExecutor(max_workers=6) as pool:
        for v in pool.map(lambda lot: traitement(consigne, lot, livrable, modele), lots):
            verdicts.update(v)
    livrable_norm = normalize(livrable)
    livrable_chiffres = unit_numbers(livrable)
    sujets_norm = " ".join(docs[n]["normes"][0] for n in role["sujets"] if n in docs)
    manques = []
    for e in elements:
        v = verdicts.get(e["id"], {})
        e["traite"] = v.get("traite", "?")
        if e["traite"] not in ("non", "partiel") or (e["traite"] == "partiel" and e.get("constat") == "conforme"):
            continue
        # Un écart ou une contradiction compte toujours ; un « absent » seulement s'il manque à la pièce vérifiée
        # ce qu'exige la référence (ce que la pièce vérifiée ajoute à la référence pèse moins).
        importance = str(e.get("importance", "moyenne")).lower()
        memes = unit_numbers(e.get("citation_reference", ""))
        if e.get("constat") == "ecart" and memes and memes == unit_numbers(e.get("citation_piece", "")):
            importance = "basse"
        a_signaler = (e.get("constat") in ("ecart", "incoherence") and importance != "basse") or (
            e.get("constat") == "absent" and e["reference"] in role["reference"] and importance == "haute"
            and dans_le_champ(e.get("section", ""), sujets_norm))
        chiffres = unit_numbers(f"{e.get('citation_reference', '')} {e.get('citation_piece', '')}")
        section = re.sub(r"(?i)^(section|article|schedule)\s+", "", str(e.get("section", ""))).strip().lower()
        deja_la = bool(chiffres) and chiffres <= livrable_chiffres and section and section in livrable_norm
        manques.append({
            "id": f"M-{len(manques) + 1:02d}", "categorie": "manque",
            "gravite": "bloquant" if a_signaler and not deja_la else "mineur",
            "section": e.get("section", ""), "sujet": e.get("sujet", ""), "constat": e.get("constat"),
            "importance": e.get("importance"),
            "reference": e["reference"], "citation_reference": e.get("citation_reference", ""),
            "piece_comparee": e.get("piece_comparee", ""), "citation_piece": e.get("citation_piece", ""),
            "detail": e.get("detail", ""), "traite": e["traite"],
            "citation_livrable": v.get("citation_livrable", "") if v.get("citation_livrable") and normalize(
                v.get("citation_livrable", "")) in livrable_norm else "",
            "a_faire": v.get("manque", ""),
        })
    stats = {"elements": len(elements), "ecartes_citation_introuvable": ecartes,
             "par_constat": {c: sum(1 for e in elements if e.get("constat") == c) for c in ("conforme", "ecart", "absent", "incoherence")},
             "non_traites": sum(1 for e in elements if e["traite"] == "non"),
             "partiels": sum(1 for e in elements if e["traite"] == "partiel")}
    return manques, stats


# ---------------------------------------------------------------- ce qui est faux, et les calculs

DATE = re.compile(r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)"
                  r"\s+\d{1,2},\s+\d{4}\b")
SECTION = re.compile(r"\b(?:Section|Schedule|Article|Exhibit|Annex)\s+(\d+(?:\.\d+)*(?:\([a-z0-9]+\))*)", re.I)


def phrase_autour(texte: str, debut: int, fin: int) -> str:
    g = max(texte.rfind("\n", 0, debut), texte.rfind(". ", 0, debut))
    d = min([p for p in (texte.find("\n", fin), texte.find(". ", fin)) if p >= 0] or [len(texte)])
    return " ".join(texte[g + 1:d + 1].split())[:400]


def introuvables(livrable: str, docs: dict) -> list[dict]:
    """Chiffres, dates et sections du livrable que le script ne retrouve dans aucune pièce."""
    tout = " ".join(d["texte"] for d in docs.values())
    tout_norm, chiffres_pieces = normalize(tout), unit_numbers(tout)
    nus = {f"{float(x.replace(',', '')):.12g}" for x in re.findall(r"\d[\d,]*(?:\.\d+)?", tout) if x.replace(",", "").replace(".", "")}
    sortie, vus = [], set()
    for m in re.finditer(r"(?:[$€£]\s?\d[\d,]*(?:\.\d+)?\s*(?:million|billion|thousand|mm|bn|m|k)?\b|\d[\d,]*(?:\.\d+)?\s*"
                         r"(?:%|percent|bps|basis points|x\b|business days?|days?|months?|years?|million|billion))",
                         livrable, re.I):
        valeurs = unit_numbers(m.group(0))
        if not valeurs or valeurs & (chiffres_pieces | nus) or m.group(0) in vus:
            continue
        vus.add(m.group(0))
        sortie.append({"type": "chiffre", "valeur": m.group(0).strip(), "phrase": phrase_autour(livrable, m.start(), m.end())})
    aujourdhui = datetime.now().strftime("%B %-d, %Y")
    for m in DATE.finditer(livrable):
        if normalize(m.group(0)) not in tout_norm and m.group(0) not in vus and m.group(0) != aujourdhui:
            vus.add(m.group(0))
            sortie.append({"type": "date", "valeur": m.group(0), "phrase": phrase_autour(livrable, m.start(), m.end())})
    for m in SECTION.finditer(livrable):
        if m.group(1).lower() not in tout_norm and m.group(0) not in vus:
            vus.add(m.group(0))
            sortie.append({"type": "section", "valeur": m.group(0), "phrase": phrase_autour(livrable, m.start(), m.end())})
    return sortie


def ce_qui_est_faux(consigne: str, docs: dict, livrable: str, candidats: list[dict], modele: str) -> dict:
    pieces = "\n\n".join(f'<piece nom="{nom}">\n{d["texte"]}\n</piece>' for nom, d in docs.items())
    liste = "\n".join(f"[I-{i + 1:03d}] {c['type']} « {c['valeur']} » dans : {c['phrase']}" for i, c in enumerate(candidats[:150]))
    return mistral(SYSTEME, f"""Consigne donnée à l'agent :
<consigne>{consigne}</consigne>

Pièces du dossier :
{pieces}

Livrable de l'agent :
<livrable>
{livrable}
</livrable>

Travail 1. Un script n'a retrouvé dans aucune pièce les valeurs suivantes du livrable. Pour chacune, dis si
c'est "calcul" (dérivée des pièces par un calcul : donne la formule avec les seuls chiffres des pièces, par
ex. "8300000 - 7900000"), "reformulation" (même valeur écrite autrement, ou référence légitime), "faux" (contredite
par les pièces : cite l'extrait qui le prouve), ou "non_source" (ni dans les pièces ni calculable).
{liste or '(aucune)'}

Travail 2. Relève toute autre affirmation du livrable contredite par les pièces : mauvais montant, mauvaise
date, mauvaise partie, mauvaise section, mauvais délai, citation déformée, document attribué à tort, ou
conclusion inverse de ce que disent les pièces. Seulement ce qui est faux de façon vérifiable, pas les avis.

Travail 3. Relève les calculs (sommes, différences, pourcentages, multiples, écarts à un seuil ou à une cible)
présents dans le livrable ou dans les pièces, avec leurs opérandes, et ceux qu'une pièce de référence impose
(par ex. un montant cible ou une définition contractuelle appliquée aux chiffres d'une autre pièce). Donne la
formule avec des chiffres seulement : un script les refera.

Réponds : {{"introuvables": [{{"id": "I-001", "verdict": "calcul|reformulation|faux|non_source",
"formule": "", "citation_piece": "copie mot pour mot d'une pièce", "piece": "nom exact", "correction": "en français"}}],
"faux": [{{"citation_livrable": "copie mot pour mot du livrable", "citation_piece": "copie mot pour mot d'une pièce",
"piece": "nom exact", "explication": "en français : ce qui est faux et la bonne valeur", "correction": "en français"}}],
"calculs": [{{"description": "en français", "ou": "livrable ou nom de la pièce", "citation": "copie mot pour mot de l'endroit où le résultat est annoncé",
"formule": "12700000 - 4800000", "resultat_annonce": "7,900,000"}}]}}""", modele, 20000)


def contexte(citation: str, doc: dict) -> str:
    """Le passage de la pièce autour de la citation (le bloc qui la contient et ses voisins)."""
    liste = [b for _, b in blocks(doc["texte"])]
    cible = normalize(citation)
    for i, b in enumerate(liste):
        if cible[:60] in normalize(b):
            return "\n\n".join(liste[max(0, i - 1):i + 2])
    return citation


def valeurs(texte: str) -> set[str]:
    return unit_numbers(texte) | {normalize(d) for d in DATE.findall(texte)} | {m.lower() for m in SECTION.findall(texte)}


def rapporte_fidelement(citation_livrable: str, docs: dict) -> str | None:
    """La pièce dont un même passage porte toutes les valeurs de la phrase du livrable : le livrable la rapporte
    (souvent pour la critiquer, « la checklist dit 5 % »), ce n'est pas une erreur du livrable."""
    cherche = valeurs(citation_livrable)
    if not cherche:
        return None
    for nom, d in docs.items():
        if cherche <= valeurs(d["texte"]):
            return nom
    return None


def confirmer(item: dict, docs: dict, consigne: str, modele: str) -> dict:
    """Un second regard, ciblé, sur une affirmation que le relecteur dit fausse : on ne garde que le contredit net."""
    cherche = valeurs(item["citation_livrable"])
    ailleurs = [f'<piece nom="{n}">{b}</piece>' for n, d in docs.items() if n != item["piece"]
                for _, b in blocks(d["texte"]) if cherche and len(cherche & valeurs(b)) >= max(1, len(cherche) // 2)][:4]
    r = mistral(SYSTEME, f"""Consigne de la tâche : {consigne[:600]}

Affirmation du livrable :
<livrable>{item['citation_livrable']}</livrable>

Passage de la pièce « {item['piece']} » :
<piece>{contexte(item['citation_piece'], docs[item['piece']])}</piece>

Passages d'autres pièces qui portent les mêmes valeurs :
{chr(10).join(ailleurs) or '(aucun)'}

Le livrable peut rapporter ce que dit une pièce pour la critiquer (« la checklist indique 5 % alors que le contrat
dit 4 % ») : ce n'est faux que si la pièce qu'il cite ne dit pas cela. L'affirmation du livrable est-elle
CONTREDITE par les pièces (valeur, date, partie, section, délai ou portée différente de ce que dit la pièce à
laquelle le livrable l'attribue) ? Réponds false si le livrable rapporte fidèlement une pièce, s'il ne fait
que compléter ou interpréter, ou si le doute existe.
Réponds : {{"contredit": true|false, "valeur_livrable": "", "valeur_piece": "", "explication": "en français, une phrase"}}""",
                modele, 1500)
    return r


# ---------------------------------------------------------------- la revue

def verifier(dossier: Path, modele: str = MODELE, tour: int | None = None) -> dict:
    debut = time.time()
    JETONS.update(entree=0, sortie=0, appels=0)
    ECARTES.clear()
    dossier = dossier.resolve()
    run_id = dossier.relative_to((banc.LAB / "results").resolve()).as_posix() if (banc.LAB / "results").resolve() in dossier.parents else ""
    nom_tache = "/".join(run_id.split("/")[:2]) if run_id else ""
    consigne, documents = consigne_et_pieces(nom_tache)
    docs = {f.relative_to(documents).as_posix(): lire(f) for f in sorted(documents.rglob("*")) if f.is_file()}
    fichiers = sorted(f for f in (dossier / "output").iterdir() if f.is_file())
    if not fichiers:
        sys.exit(f"aucun livrable dans {dossier / 'output'}")
    livrables = {f.name: lire(f) for f in fichiers}
    livrable = "\n\n".join(f"=== {n} ===\n{d['texte']}" for n, d in livrables.items())

    # Calculs refaits par script : tableurs et opérations écrites, dans les pièces et dans le livrable.
    calculs_script = []
    for f in [*sorted(documents.rglob("*.xlsx")), *[f for f in fichiers if f.suffix == ".xlsx"]]:
        calculs_script += totaux_tableurs(f)
    for nom, d in [*docs.items(), *livrables.items()]:
        calculs_script += operations_ecrites(nom, d["texte"])
        if nom in docs and not nom.endswith((".xlsx", ".xls", ".csv")):  # les tableaux d'analyse du livrable ne sont pas des sommes
            calculs_script += totaux_texte(nom, d["texte"])

    role = roles(consigne, docs, modele)
    candidats = introuvables(livrable, docs)
    with ThreadPoolExecutor(max_workers=2) as pool:
        tache_manques = pool.submit(ce_qui_manque, consigne, docs, livrable, role, modele)
        tache_faux = pool.submit(ce_qui_est_faux, consigne, docs, livrable, candidats, modele)
        manques, stats = tache_manques.result()
        faux_brut = tache_faux.result()

    problemes, ecartes, a_confirmer = list(manques), 0, []
    # Ce qui est faux : affirmations contredites, extraits vérifiés des deux côtés, puis confirmées.
    for f in faux_brut.get("faux") or []:
        piece = ou_est(f.get("citation_piece", ""), docs, nom_proche(f.get("piece", ""), docs))
        dans_livrable = any(retrouve(f.get("citation_livrable", ""), d) for d in livrables.values())
        if not (piece and dans_livrable):
            ecartes += 1
            ECARTES.append({"raison": "faux : citation introuvable", **f})
            continue
        a_confirmer.append({"categorie": "faux", "gravite": "bloquant", "citation_livrable": f["citation_livrable"],
                            "piece": piece, "citation_piece": f["citation_piece"], "detail": f.get("explication", ""),
                            "a_faire": f.get("correction", "")})
    # Valeurs introuvables : contredites (bloquant), calculées (recalcul), non sourcées (mineur).
    par_id = {f"I-{i + 1:03d}": c for i, c in enumerate(candidats[:150])}
    for v in faux_brut.get("introuvables") or []:
        c = par_id.get(v.get("id"))
        if not c:
            continue
        verdict = v.get("verdict")
        if verdict == "faux":
            piece = ou_est(v.get("citation_piece", ""), docs, nom_proche(v.get("piece", ""), docs))
            if piece:
                a_confirmer.append({"categorie": "faux", "gravite": "bloquant", "citation_livrable": c["phrase"],
                                    "valeur": c["valeur"], "piece": piece, "citation_piece": v["citation_piece"],
                                    "detail": f"« {c['valeur']} » est contredit par les pièces.", "a_faire": v.get("correction", "")})
            else:
                ecartes += 1
                ECARTES.append({"raison": "introuvable dit faux : citation de la pièce introuvable", **c, **v})
        elif verdict == "calcul":
            calcul = evaluer(v.get("formule", ""))
            annonce = next(iter(unit_numbers(c["valeur"])), None)
            if calcul is not None and annonce is not None and not arrondi_compatible(float(annonce), calcul):
                net = plausible(float(annonce), calcul)
                problemes.append({"categorie": "calcul", "gravite": "bloquant" if net else "mineur", "citation_livrable": c["phrase"],
                                  "valeur": c["valeur"], "formule": v.get("formule"), "recalcule": calcul,
                                  "detail": f"Le livrable annonce {c['valeur']} ; refait par script, {v.get('formule')} = {calcul:,.2f}."
                                  + ("" if net else " Écart trop grand pour une simple erreur : la formule du relecteur est peut-être fausse."),
                                  "a_faire": "Refaire ce calcul à partir des pièces et corriger le chiffre partout où il apparaît."})
        elif verdict == "non_source":
            problemes.append({"categorie": "faux", "gravite": "mineur", "citation_livrable": c["phrase"], "valeur": c["valeur"],
                              "detail": f"« {c['valeur']} » ne figure dans aucune pièce et n'en découle par aucun calcul identifié.",
                              "a_faire": "Sourcer ce chiffre (pièce et section) ou le retirer."})
    gardes = []
    for f in a_confirmer:
        source = rapporte_fidelement(f["citation_livrable"], docs)
        if not valeurs(f["citation_livrable"]) or not (valeurs(f["citation_piece"]) - valeurs(f["citation_livrable"])):
            ecartes += 1  # « faux » ne vaut que pour une valeur (montant, date, %, section) que la pièce donne autrement
            ECARTES.append({"raison": "faux écarté : pas de valeur différente dans l'extrait de la pièce", **f})
        elif source:
            ecartes += 1
            ECARTES.append({"raison": f"faux écarté : le livrable rapporte fidèlement {source}", **f})
        else:
            gardes.append(f)
    a_confirmer = gardes
    with ThreadPoolExecutor(max_workers=6) as pool:
        avis = list(pool.map(lambda f: confirmer(f, docs, consigne, modele), a_confirmer))
    for f, a in zip(a_confirmer, avis):
        if a.get("contredit") is True:
            f["detail"] = a.get("explication") or f["detail"]
            if a.get("valeur_piece"):
                f["detail"] += f" (livrable : {a.get('valeur_livrable', '?')} ; pièce : {a['valeur_piece']})"
            problemes.append(f)
        else:
            ecartes += 1
            ECARTES.append({"raison": "faux non confirmé au second regard", **f, "avis": a})
    # Calculs repérés par le relecteur, refaits par script.
    for c in faux_brut.get("calculs") or []:
        calcul = evaluer(c.get("formule", ""))
        annonce = nombre(c.get("resultat_annonce"))
        if calcul is None or annonce is None or arrondi_compatible(annonce, calcul):
            continue
        lieu = c.get("ou", "")
        dans = (any(retrouve(c.get("citation", ""), d) for d in livrables.values()) if "livrable" in lieu.lower()
                else bool(ou_est(c.get("citation", ""), docs, nom_proche(lieu, docs))))
        if not dans:
            ecartes += 1
            continue
        problemes.append({"categorie": "calcul", "gravite": "bloquant" if plausible(annonce, calcul) else "mineur",
                          "ou": lieu, "citation": c["citation"],
                          "formule": c["formule"], "recalcule": calcul, "annonce": annonce,
                          "detail": f"{c.get('description', '')} : annoncé {annonce:,.2f}, refait par script {c['formule']} = {calcul:,.2f}.",
                          "a_faire": ("Signaler cette erreur de la pièce dans le livrable, avec le bon montant et sa conséquence."
                                      if "livrable" not in lieu.lower() else "Corriger ce calcul dans le livrable.")})
    regroupes: dict = {}
    for c in calculs_script:
        cle = (c.get("libelle", c.get("extrait", "")).lower(), c["annonce"], round(c["recalcule"], 2))
        if cle in regroupes:
            regroupes[cle]["aussi"].append(f"{c['piece']}" + (f", {c['ou']}" if c.get("ou") else ""))
        else:
            regroupes[cle] = {**c, "aussi": []}
    for c in regroupes.values():
        dans_livrable = c["piece"] in livrables
        detail = (f"{c['piece']}, {c['ou']}, ligne « {c['libelle']} » : annoncé {c['annonce']:,.0f} ; refait par script : "
                  + " + ".join(f"{v:,.0f} ({l})" for l, v in c["operandes"]) + f" = {c['recalcule']:,.0f}" if "operandes" in c else
                  f"{c['piece']} : « {c['extrait']} » ; refait par script : {c['recalcule']:,.2f}")
        if c["aussi"]:
            detail += f". Même erreur dans : {'; '.join(c['aussi'])}"
        problemes.append({"categorie": "calcul", "gravite": "bloquant", "ou": c["piece"], "detail": detail, "script": c,
                          "a_faire": "Corriger ce calcul dans le livrable." if dans_livrable else
                          "Signaler cette erreur de calcul de la pièce dans le livrable, avec le bon montant et sa conséquence."})
    # Une erreur de calcul d'une pièce que le livrable signale déjà n'est plus bloquante.
    livrable_chiffres = unit_numbers(livrable) | {f"{float(x.replace(',', '')):.12g}" for x in re.findall(r"\d[\d,]*(?:\.\d+)?", livrable)}
    for p in problemes:
        if p["categorie"] == "calcul" and p.get("ou") not in livrables and "livrable" not in str(p.get("ou", "")).lower():
            juste = p.get("recalcule") if p.get("recalcule") is not None else p.get("script", {}).get("recalcule")
            if juste is not None and f"{abs(float(juste)):.12g}" in livrable_chiffres:
                p["gravite"] = "mineur"
                p["detail"] += " (le livrable semble déjà donner le bon montant : vérifier qu'il signale l'erreur)"

    ordre = {"bloquant": 0, "mineur": 1}
    rang = {"haute": 0, "moyenne": 1, "basse": 2}
    problemes.sort(key=lambda p: (ordre[p["gravite"]], ["manque", "faux", "calcul"].index(p["categorie"]),
                                  rang.get(str(p.get("importance", "haute")).lower(), 1)))
    compte = {c: sum(1 for p in problemes if p["categorie"] == c and p["gravite"] == "bloquant") for c in ("manque", "faux", "calcul")}
    for i, p in enumerate(problemes):
        p["id"] = f"{p['categorie'][0].upper()}-{i + 1:02d}"
    revue = {
        "at": datetime.now().isoformat(timespec="seconds"), "run_id": run_id, "tache": nom_tache, "tour": tour,
        "modele": modele, "livrables": list(livrables), "roles": role, "stats_concordance": stats,
        "valeurs_introuvables": len(candidats), "ecartes_citation_introuvable": stats["ecartes_citation_introuvable"] + ecartes,
        "bloquants": sum(compte.values()), "bloquants_par_categorie": compte,
        "mineurs": sum(1 for p in problemes if p["gravite"] == "mineur"), "problemes": problemes,
        "jetons": dict(JETONS), "cout_usd": cout(modele), "secondes": round(time.time() - debut, 1),
        "ecartes_detail": ECARTES[:200],
    }
    (dossier / "revue-visa.json").write_text(json.dumps(revue, indent=2, ensure_ascii=False))
    (dossier / "revue-visa.md").write_text(en_markdown(revue))
    (dossier / "revue-legora.txt").write_text(pour_legora(revue))
    return revue


def cout(modele: str) -> float:
    entree, sortie = next((t for k, t in TARIFS.items() if modele.startswith(k)), TARIFS["mistral-large"])
    return round(JETONS["entree"] / 1e6 * entree + JETONS["sortie"] / 1e6 * sortie, 4)


def consigne_et_pieces(nom_tache: str) -> tuple[str, Path]:
    """La consigne (sans la grille) et le dossier des pièces. Même source que l'agent : externe.tache()."""
    import externe

    t = externe.tache(nom_tache)
    return t["consigne"], t["documents"]


def citer(texte: str) -> str:
    return "« " + " ".join(str(texte).split()) + " »"


def bloc(p: dict) -> str:
    lignes = []
    if p["categorie"] == "manque":
        lignes.append(f"### {p['id']} [{p['gravite']}{', importance ' + p['importance'] if p.get('importance') else ''}] "
                      f"{p.get('section', '')} — {p.get('sujet', '')}")
        lignes.append(f"- Pièce de référence ({p['reference']}, {p.get('section', '')}) : {citer(p['citation_reference'])}")
        if p.get("citation_piece"):
            lignes.append(f"- Pièce comparée ({p['piece_comparee']}) : {citer(p['citation_piece'])}")
        lignes.append(f"- Constat du vérificateur : {p.get('constat')} — {p.get('detail', '')}")
        etat = "absent du livrable" if p["traite"] == "non" else "traité en partie"
        lignes.append(f"- Dans le livrable : {etat}" + (f", ici : {citer(p['citation_livrable'])}" if p.get("citation_livrable") else ""))
    elif p["categorie"] == "faux":
        lignes.append(f"### {p['id']} [{p['gravite']}] {p.get('valeur') or 'Affirmation contredite par les pièces'}")
        lignes.append(f"- Livrable : {citer(p['citation_livrable'])}")
        if p.get("citation_piece"):
            lignes.append(f"- Pièce ({p['piece']}) : {citer(p['citation_piece'])}")
        lignes.append(f"- Problème : {p.get('detail', '')}")
    else:
        lignes.append(f"### {p['id']} [{p['gravite']}] Calcul à reprendre ({p.get('ou', '')})")
        if p.get("citation"):
            lignes.append(f"- Endroit : {citer(p['citation'])}")
        if p.get("citation_livrable"):
            lignes.append(f"- Livrable : {citer(p['citation_livrable'])}")
        lignes.append(f"- Refait par script : {p.get('detail', '')}")
    if p.get("a_faire"):
        lignes.append(f"- À faire : {p['a_faire']}")
    return "\n".join(lignes)


def en_markdown(r: dict) -> str:
    c = r["bloquants_par_categorie"]
    tete = [f"# Revue Visa — {', '.join(r['livrables'])}" + (f" (tour {r['tour']})" if r.get("tour") is not None else ""), "",
            "Un vérificateur indépendant (Mistral) a relu le livrable contre les pièces du dossier. Chaque point cite les pièces "
            "mot pour mot (extraits vérifiés par script) et chaque calcul a été refait par script. Le vérificateur peut se "
            "tromper : vérifie chaque point dans les pièces avant de corriger, et ignore un point qui ne tient pas.", "",
            f"**Problèmes bloquants : {r['bloquants']}** (manques : {c['manque']}, faux : {c['faux']}, calculs : {c['calcul']}). "
            f"Points mineurs : {r['mineurs']}.", "",
            f"Pièce(s) de référence retenue(s) : {', '.join(r['roles']['reference'])}. Éléments relevés dans la référence : "
            f"{r['stats_concordance']['elements']}.", ""]
    sections = [("Ce qui manque", "manque"), ("Ce qui est faux", "faux"), ("Calculs", "calcul")]
    corps = []
    for titre, cat in sections:
        liste = [p for p in r["problemes"] if p["categorie"] == cat and p["gravite"] == "bloquant"]
        if liste:
            corps += [f"## {titre}", ""] + [bloc(p) + "\n" for p in liste]
    mineurs = [p for p in r["problemes"] if p["gravite"] == "mineur"]
    if mineurs:
        corps += ["## Points mineurs (à traiter si c'est fondé)", ""] + [bloc(p) + "\n" for p in mineurs[:25]]
        if len(mineurs) > 25:
            corps.append(f"(et {len(mineurs) - 25} autres points mineurs, dans revue-visa.json)")
    if not r["problemes"]:
        corps.append("Aucun problème trouvé.")
    return "\n".join(tete + corps) + "\n"


def pour_legora(r: dict, maximum: int = 25) -> str:
    """Le texte court à coller dans Legora, à la suite de la conversation qui a produit le livrable."""
    lignes = [f"Un vérificateur indépendant a relu ton livrable ({', '.join(r['livrables'])}). Corrige le livrable en "
              "conséquence : ajoute ce qui manque, corrige ce qui est faux, garde ce qui est juste. Même nom de fichier, "
              "même format, même langue. Vérifie chaque point dans les pièces avant de corriger, et ignore un point qui ne tient pas.",
              "", f"Problèmes à corriger ({r['bloquants']}) :"]
    for n, p in enumerate([p for p in r["problemes"] if p["gravite"] == "bloquant"][:maximum], 1):
        if p["categorie"] == "manque":
            texte = (f"{p.get('section', '')} — {p.get('sujet', '')} : la référence dit {citer(p['citation_reference'])}"
                     + (f" ; {p['piece_comparee']} dit {citer(p['citation_piece'])}" if p.get("citation_piece") else "")
                     + f". {p.get('detail', '')} Le livrable {'ne le traite pas' if p['traite'] == 'non' else 'ne le traite qu’en partie'}."
                     + (f" À faire : {p['a_faire']}" if p.get("a_faire") else ""))
        elif p["categorie"] == "faux":
            texte = (f"Faux : le livrable dit {citer(p['citation_livrable'])}"
                     + (f" alors que {p['piece']} dit {citer(p['citation_piece'])}" if p.get("citation_piece") else "")
                     + f". {p.get('detail', '')}")
        else:
            texte = f"Calcul : {p.get('detail', '')} {p.get('a_faire', '')}"
        lignes.append(f"{n}. {' '.join(texte.split())}")
    return "\n".join(lignes) + "\n"


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("lancement", help="dossier du lancement (avec output/) ou run_id")
    p.add_argument("--modele", default=MODELE, help="mistral-large-latest ou mistral-medium-3.5")
    a = p.parse_args()
    dossier = Path(a.lancement)
    if not dossier.exists():
        dossier = banc.LAB / "results" / a.lancement.strip("/").removeprefix("results/")
    r = verifier(dossier, a.modele)
    c = r["bloquants_par_categorie"]
    print(f"{r['bloquants']} bloquants (manques {c['manque']}, faux {c['faux']}, calculs {c['calcul']}), {r['mineurs']} mineurs ; "
          f"{r['stats_concordance']['elements']} éléments relevés, {r['ecartes_citation_introuvable']} écartés (citation introuvable) ; "
          f"{r['secondes']} s, {r['jetons']['entree']:,} jetons en entrée, ~{r['cout_usd']} $ — {dossier / 'revue-visa.md'}")


if __name__ == "__main__":
    main()
