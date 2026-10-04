"""Exports du même registre vers HTML, Word et Excel."""
import json
import textwrap
from pathlib import Path


def tables(data):
    risks = [[f['id'], f['workstream'], f['topic'], f['priority'], f['statement'], f['impact'], f['recommendation'], f['validation']] for f in data['findings']]
    docs = [[d['id'], d['path'], d['extraction'], d['decision'], d['decision_reason'], 'Oui (déclaré)' if d['consulted'] else 'Non', d['reason'], d['parent'] or '', d['signature']] for d in data['documents']]
    questions = [[f['id'], f['question'], f.get('seller_response', ''), f.get('response_source', ''), f.get('owner', 'À attribuer'), f.get('due_date', ''), 'À rapprocher des pièces' if f.get('seller_response') else 'En attente'] for f in data['findings'] if f.get('question')]
    evidence = [[f['id'], e.get('path', ''), e.get('location', ''), e.get('clause', ''), e.get('quote', ''), 'Ancrage vérifié' if e['verified'] else e.get('error', ''), f['selection_reason']] for f in data['findings'] for e in f['evidence']]
    return [('Risques', ['ID', 'Volet', 'Sujet', 'Priorité', 'Constat', 'Impact', 'Suite proposée', 'Validation'], risks),
            ('Pièces', ['ID', 'Chemin', 'Extraction', 'Décision', 'Motif', 'Consultation', 'Limite', 'Contrat parent', 'Signature'], docs),
            ('Questions vendeur', ['Constat', 'Question', 'Réponse vendeur', 'Référence réponse', 'Responsable', 'Échéance', 'État'], questions),
            ('Preuves', ['Constat', 'Pièce', 'Localisation', 'Clause', 'Extrait exact', 'Contrôle technique', 'Motif du choix'], evidence)]


def export_html(data, out):
    template = Path(__file__).with_name('review.html').read_text()
    (out / 'revue.html').write_text(template.replace('__DATA__', json.dumps(data, ensure_ascii=False).replace('<', '\\u003c')), encoding='utf-8')


def export_office(data, out):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    report = Document()
    for border in report.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = report.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.6)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = report.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    report.styles['Normal'].font.size = Pt(10)
    report.styles['Normal'].paragraph_format.space_after = Pt(6)
    section.header.paragraphs[0].text = 'CONFIDENTIEL  ·  DOCUMENT DE TRAVAIL'
    footer = section.footer.paragraphs[0]
    footer.text = 'Due diligence  ·  '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    def table(headers, rows, widths):
        t = report.add_table(rows=1, cols=len(headers))
        t.autofit = False
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
        t._tbl.tblPr.append(borders)
        repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for i, values in enumerate([headers] + rows):
            cells = t.rows[0].cells if i == 0 else t.add_row().cells
            no_split = OxmlElement('w:cantSplit')
            t.rows[i]._tr.get_or_add_trPr().append(no_split)
            for cell, value, width in zip(cells, values, widths):
                cell.width = Cm(width); cell.text = str(value)
                margins = OxmlElement('w:tcMar')
                for side in ['top', 'left', 'bottom', 'right']:
                    el = OxmlElement('w:' + side); el.set(qn('w:w'), '85'); el.set(qn('w:type'), 'dxa'); margins.append(el)
                cell._tc.get_or_add_tcPr().append(margins)
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '183F45' if i == 0 else 'F1F5F5' if i % 2 else 'FFFFFF'); cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if i == 0:
                            run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
        report.add_paragraph()

    report.add_paragraph('Rapport de due diligence', 'Title')
    report.add_paragraph(data['mission']['title'])
    report.add_paragraph(f"Destinataire : {data['mission']['client']}\nSituation documentaire au {data['mission']['as_of']}")
    report.add_paragraph('Brouillon à valider' if data['draft'] else 'Constats validés dans le périmètre déclaré', 'Heading 1')
    report.add_paragraph(data['mission']['scope'])
    report.add_paragraph(f"{len(data['findings'])} constats ; {len(data['documents'])} pièces inventoriées ; {sum(d['consulted'] for d in data['documents'])} pièces déclarées consultées. L’ancrage technique d’une citation ne valide ni l’interprétation juridique ni l’exhaustivité de la revue.")
    for limitation in data['limitations']:
        report.add_paragraph(limitation, 'List Bullet')
    report.add_heading('Synthèse des risques', 1)
    table(['ID et priorité', 'Constat', 'Suite proposée'], [[f"{f['id']}\n{f['priority']}\n{f['validation']}", f['statement'], f['recommendation']] for f in data['findings']], [3, 8, 6.8])
    report.add_page_break()
    report.add_heading('Fiches de vérification', 1)
    for index, f in enumerate(data['findings']):
        if index:
            report.add_page_break()
        report.add_heading(f"{f['id']} {f['topic']}", 2)
        report.add_paragraph(f"{f['workstream']} · {f['priority']} · {f['validation']}")
        report.add_paragraph(f['statement'])
        report.add_paragraph('Impact : ' + f['impact'])
        report.add_paragraph('Suite proposée : ' + f['recommendation'])
        report.add_paragraph('Pièces consultées (déclaration) : ' + ', '.join(f['consulted']))
        report.add_paragraph('Choix de la source : ' + f['selection_reason'])
        table(['Source et localisation', 'Extrait exact et contrôle'], [[f"{e.get('path', '')}\n{e.get('location', '')}\n{e.get('clause', '')}", e.get('quote', '') + '\n' + ('Ancrage vérifié' if e['verified'] else 'PREUVE À REPRENDRE : ' + e.get('error', ''))] for e in f['evidence']], [6, 11.8])
        if f.get('links'):
            report.add_paragraph('Constats liés : ' + ', '.join(f['links']))
        if f.get('reviewer'):
            report.add_paragraph(f"Relecteur déclaré : {f['reviewer']} · {f['reviewed_at']}")
    report.add_page_break()
    report.add_heading('Questions au vendeur', 1)
    table(['ID', 'Question', 'Réponse et suivi'], [[f['id'], f['question'], f.get('seller_response') or 'En attente de réponse'] for f in data['findings'] if f.get('question')], [2, 9, 6.8])
    report.add_heading('Registre des pièces et limites', 1)
    table(['Pièce', 'Traitement', 'Motif et limites'], [[d['path'], f"{d['extraction']}\n{d['decision']}\nConsultation : {'déclarée' if d['consulted'] else 'non déclarée'}", '\n'.join(x for x in [d['decision_reason'], d['reason'], 'Signature non vérifiée', 'Contrat parent : ' + d['parent'] if d['parent'] else ''] if x)] for d in data['documents']], [6, 4, 7.8])
    report.add_heading('Nouveaux versements', 1)
    if data['changes']:
        table(['Pièce', 'Changement'], [[c['path'], c['status']] for c in data['changes']], [13, 4.8])
    else:
        report.add_paragraph('Aucun écart détecté avec le registre précédent, ou premier inventaire.')
    report.save(out / 'rapport.docx')

    book = Workbook(); book.remove(book.active)
    for name, headers, rows in tables(data):
        sheet = book.create_sheet(name)
        for values in [headers] + rows:
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # empêcher toute formule issue des pièces
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.font = Font(name='Calibri', size=11)
        sheet.freeze_panes = 'B2'
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='183F45'); cell.font = Font(color='FFFFFF', bold=True)
        widths = {
            'Risques': [14, 24, 44, 16, 68, 68, 68, 18],
            'Pièces': [20, 60, 18, 18, 62, 24, 58, 22, 22],
            'Questions vendeur': [14, 78, 65, 40, 20, 18, 30],
            'Preuves': [14, 48, 28, 28, 90, 28, 70],
        }[name]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.row_dimensions[1].height = 30
        for index in range(2, sheet.max_row + 1):
            line_counts = []
            for cell, width in zip(sheet[index], widths):
                lines = str(cell.value or '').split('\n')
                line_counts.append(sum(max(1, len(textwrap.wrap(line, width=max(8, int(width * .85))))) for line in lines))
            sheet.row_dimensions[index].height = min(409, max(55, max(line_counts) * 16 + 12))
        if rows:
            t = Table(displayName='Table' + str(len(book.worksheets)), ref=sheet.dimensions)
            t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(t)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'
    limits = book.create_sheet('Périmètre')
    for row in [('Mission', data['mission']['title']), ('Périmètre', data['mission']['scope']), ('Date de référence', data['mission']['as_of']), ('État', 'Brouillon à valider' if data['draft'] else 'Constats validés'), *[('Limite', x) for x in data['limitations']], *[('Versement', c['path'] + ' : ' + c['status']) for c in data['changes']]]:
        limits.append(row)
        for cell in limits[limits.max_row]:
            cell.data_type = 's'; cell.alignment = Alignment(wrap_text=True, vertical='top')
        limits.row_dimensions[limits.max_row].height = 60
    limits.column_dimensions['A'].width = 25; limits.column_dimensions['B'].width = 100
    book.save(out / 'tableaux.xlsx')
