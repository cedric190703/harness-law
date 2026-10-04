import copy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from due_diligence.review import build, inventory
from due_diligence.exports import export_html, export_office


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.room = self.root / 'pieces'; self.room.mkdir()
        self.path = self.room / 'contrat.txt'
        self.path.write_text('Section 1\nPréavis de soixante jours.\nFin du contrat.\n')
        self.sha = hashlib.sha256(self.path.read_bytes()).hexdigest()
        self.case = dict(title='Audit', client='Acquéreur', scope='Contrats', as_of='2025-01-08', findings=[dict(
            id='DD-1', topic='Préavis', workstream='Commercial', priority='Élevée', statement='Préavis de 60 jours.',
            impact='Revenus exposés.', recommendation='Obtenir le contrat signé.', selection_reason='Clause expresse.',
            consulted=['contrat.txt'], evidence=[dict(path='contrat.txt', sha256=self.sha, clause='Section 1',
            line_start=2, line_end=2, quote='Préavis de soixante jours.')])])

    def test_quote_is_located_and_extraction_does_not_imply_consultation(self):
        (self.room / 'autre.txt').write_text('Autre pièce.')
        data = build(self.room, self.case)
        self.assertTrue(data['findings'][0]['evidence'][0]['verified'])
        self.assertFalse(next(d for d in data['documents'] if d['path']=='autre.txt')['consulted'])
        self.assertTrue(data['draft'])

    def test_false_quote_and_wrong_location_are_rejected(self):
        for patch_value in [{'quote':'Préavis de trente jours.'}, {'line_start':1, 'line_end':1}]:
            case = copy.deepcopy(self.case); case['findings'][0]['evidence'][0].update(patch_value)
            data = build(self.room, case)
            self.assertTrue(data['findings'][0]['errors'])
            self.assertEqual(data['findings'][0]['validation'], 'à reprendre')

    def test_modified_source_invalidates_validated_finding(self):
        case = copy.deepcopy(self.case)
        case['findings'][0].update(validation='validé', reviewer='Juriste', reviewed_at='2025-01-08')
        old = build(self.room, case)
        self.assertEqual(old['findings'][0]['validation'], 'validé')
        self.path.write_text('Section 1\nPréavis de trente jours.\n')
        current = build(self.room, case, old)
        self.assertEqual(current['changes'][0]['status'], 'modifiée')
        self.assertEqual(current['findings'][0]['validation'], 'à reprendre')

    def test_new_document_requires_recheck(self):
        case = copy.deepcopy(self.case)
        case['findings'][0].update(validation='validé', reviewer='Juriste', reviewed_at='2025-01-08')
        old = build(self.room, case)
        (self.room / 'avenant.txt').write_text('Nouveau préavis.')
        new = build(self.room, case, old)
        self.assertEqual(new['findings'][0]['validation'], 'à reprendre')

    def test_duplicate_homonym_and_unreadable_are_visible(self):
        sub = self.room / 'sous-dossier'; sub.mkdir()
        (sub / 'contrat.txt').write_bytes(self.path.read_bytes())
        (self.room / 'image.bin').write_bytes(b'\x00\x01')
        docs = inventory(self.room)
        self.assertEqual(len({d['id'] for d in docs}), 3)
        self.assertEqual(sum(d['duplicate_of'] is not None for d in docs), 1)
        self.assertEqual(next(d for d in docs if d['path']=='image.bin')['extraction'], 'illisible')

    def test_partial_pdf_and_symlink_are_not_silently_read(self):
        fake = self.room / 'scan.pdf'; fake.write_bytes(b'fake')
        with patch('due_diligence.review.extract', return_value=(['texte', ''], 'pages')):
            docs = inventory(self.room)
        self.assertEqual(docs[0]['unread_units'], [2])
        self.assertEqual(docs[0]['extraction'], 'partielle')
        (self.room / 'link.txt').symlink_to(self.path)
        self.assertEqual(next(d for d in inventory(self.room) if d['path']=='link.txt')['extraction'], 'illisible')

    def test_excluded_source_cannot_support_finding(self):
        self.case['documents'] = {'contrat.txt':dict(decision='exclue', reason='Projet non signé confirmé', sha256=self.sha)}
        self.assertTrue(build(self.room, self.case)['findings'][0]['errors'])

    def test_validation_requires_reviewer(self):
        self.case['findings'][0]['validation']='validé'
        with self.assertRaises(ValueError): build(self.room, self.case)

    def test_repeated_ids_and_broken_links_fail(self):
        repeated=copy.deepcopy(self.case); repeated['findings'] *= 2
        with self.assertRaises(ValueError): build(self.room, repeated)
        self.case['findings'][0]['links']=['unknown']
        with self.assertRaises(ValueError): build(self.room, self.case)

    def test_html_escapes_script_terminators(self):
        self.case['title']='</script><script>alert(1)</script>'
        data=build(self.room, self.case); export_html(data, self.root)
        self.assertNotIn('</script><script>alert(1)', (self.root/'revue.html').read_text())

    def test_exports_keep_evidence_and_excel_cells_are_text(self):
        from docx import Document
        from openpyxl import load_workbook
        self.case['findings'][0]['statement']='=HYPERLINK("https://invalid")'
        data=build(self.room, self.case); export_office(data,self.root)
        book=load_workbook(self.root/'tableaux.xlsx')
        self.assertEqual(book['Risques']['E2'].data_type,'s')
        self.assertIn('Préavis de soixante jours.',book['Preuves']['E2'].value)
        self.assertEqual(book['Risques'].freeze_panes,'B2')
        doc=Document(self.root/'rapport.docx')
        content='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn('Préavis de soixante jours.',content)
        self.assertIn('l. 2–2',content)

if __name__ == '__main__': unittest.main()
