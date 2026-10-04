#!/usr/bin/env python3
"""Inventaire de la data room : texte de chaque pièce, empreinte, signaux, registre.

  python inventaire.py [dossier_des_pièces]          # défaut : $DOCUMENTS_DIR

Écrit $WORKSPACE_DIR/dd/texte/<rubrique>/<pièce>.txt (repères « === Page n === » ou
« === Feuille x === », lignes « L12: … » pour les tableurs), dd/registre.jsonl (une ligne par
pièce, statut à remplir) et dd/inventaire.md. Relancer conserve les statuts et fiches déjà saisis.
"""
import csv
import difflib
import email
import email.policy
import hashlib
import html
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path

import commun as c

SYSTEME = re.compile(r"^(\.ds_store|thumbs\.db|desktop\.ini|~\$.*|\._.*|\.~lock.*)$", re.IGNORECASE)
TABLEURS = {".xlsx", ".xlsm", ".xls", ".csv", ".ods"}
PROJET = (r"projet|draft|brouillon|non signee?|a signer|pour discussion|provisoire|v\d+|version|modele|template|"
          r"specimen|exemple")
VERSION = PROJET + r"|copie|final|finale|definitif|definitive|signe|signee"
PROJET_NOM = re.compile(rf"\b({PROJET})\b")
PROJET_DEBUT = re.compile(r"\b(projet|draft|brouillon|document de travail|pour discussion|version provisoire|"
                          r"non signee?|a signer)\b")
PROJET_TEXTE = re.compile(r"(non signee?|en attente de signature|\[signature\]|_{8,}|\.{12,}|"
                          r"sous reserve de (sa )?signature)")
SIGNATURE = re.compile(r"\bsign(e|ee|es|ees|ature|atures|ataire|ataires)\b|lu et approuve")
ACTE = re.compile(r"\b(contrat|avenant|accord|convention|bail|acte|promesse|lettre|proces verbal|pv|decision|statuts|"
                  r"attestation|certificat|protocole|pacte|engagement|cession|transaction|police)\b")
AVENANT = re.compile(r"\b(avenant|amendement|addendum|annexe|waiver|renonciation|side letter|accord modificatif|"
                     r"rectificatif|lettre complementaire)\b")
GENERIQUES = set(("avenant amendement addendum annexe waiver renonciation rectificatif contrat contrats accord accords "
                  "convention projet signe signee version final finale definitif copie lettre rapport memo note dossier "
                  "tableau liste registre etat suivi fiche document documents demande confirmation reponse modele "
                  "attestation certificat proces verbal decision draft side complementaire modificatif relatif "
                  "relative relatifs relatives avec pour dans entre").split())


def lancer(cmd: list[str]) -> str:
    return subprocess.run(cmd, capture_output=True, text=True, check=True, timeout=180).stdout


def avec_pages(pages: list[str]) -> str:
    return "\n".join(f"=== Page {i} ===\n{page.rstrip()}" for i, page in enumerate(pages, 1))


def ocr_image(image: Path) -> str:
    try:
        return lancer(["tesseract", str(image), "-", "-l", "fra+eng"])
    except Exception:
        return lancer(["tesseract", str(image), "-"])


def ocr_pdf(chemin: Path) -> list[str]:
    if not (shutil.which("tesseract") and shutil.which("pdftoppm")):
        return []
    with tempfile.TemporaryDirectory(dir=c.DD) as dossier:
        lancer(["pdftoppm", "-r", "200", "-png", str(chemin), f"{dossier}/p"])
        return [ocr_image(image) for image in sorted(Path(dossier).glob("p*.png"))]


def pages_pdf(chemin: Path) -> tuple[list[str], str]:
    def pdftotext():
        pages = lancer(["pdftotext", "-layout", "-enc", "UTF-8", str(chemin), "-"]).split("\f")
        return pages[:-1] if len(pages) > 1 and not pages[-1].strip() else pages

    def pdfplumber():
        import pdfplumber as module
        with module.open(str(chemin)) as pdf:
            return [page.extract_text() or "" for page in pdf.pages]

    def fitz():
        import fitz as module
        with module.open(str(chemin)) as pdf:
            return [page.get_text() for page in pdf]

    def pypdf():
        import pypdf as module
        return [page.extract_text() or "" for page in module.PdfReader(str(chemin)).pages]

    for essai in ([pdftotext] if shutil.which("pdftotext") else []) + [pdfplumber, fitz, pypdf]:
        try:
            pages = essai()
        except Exception:
            continue
        if any(page.strip() for page in pages):
            return pages, ""
    try:
        pages = ocr_pdf(chemin)
    except Exception:
        pages = []
    if any(page.strip() for page in pages):
        return pages, "texte obtenu par OCR : vérifier chaque extrait sur l'original"
    try:
        import fitz as module
        with module.open(str(chemin)) as pdf:
            vides = [""] * pdf.page_count
    except Exception:
        vides = []
    return vides, "aucun texte extractible (scan ?)" + ("" if shutil.which("tesseract") else ", OCR indisponible")


def nombre(v) -> str:
    v = float(v)
    return str(int(v)) if v.is_integer() else f"{v:.6f}".rstrip("0").rstrip(".")


def valeur(v, format_nombre: str = "") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VRAI" if v else "FAUX"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return nombre(v * 100) + " %" if "%" in (format_nombre or "") else nombre(v)
    return " ".join(str(v).split())


def feuilles_xlsx(chemin: Path) -> tuple[str, int]:
    import openpyxl

    valeurs = openpyxl.load_workbook(chemin, data_only=True)
    formules = openpyxl.load_workbook(chemin)
    sortie = []
    for feuille in valeurs.worksheets:
        lignes = []
        for rangee in feuille.iter_rows():
            cellules = []
            for cellule in rangee:
                v = cellule.value
                if v is None:  # formule sans valeur calculée : on montre la formule
                    f = formules[feuille.title][cellule.coordinate].value
                    v = f if isinstance(f, str) and f.startswith("=") else None
                cellules.append(valeur(v, getattr(cellule, "number_format", "")))
            while cellules and not cellules[-1]:
                cellules.pop()
            if cellules:
                lignes.append(f"L{rangee[0].row}: " + " | ".join(cellules))
        sortie.append(f"=== Feuille {feuille.title} ===\n" + "\n".join(lignes))
    return "\n".join(sortie), len(valeurs.worksheets)


def feuilles_autres(chemin: Path) -> tuple[str, int]:
    if chemin.suffix.lower() == ".csv":
        brut = chemin.read_text(encoding="utf-8", errors="replace")
        dialecte = csv.Sniffer().sniff(brut[:4000]) if brut.strip() else csv.excel
        lignes = [f"L{n}: " + " | ".join(r) for n, r in enumerate(csv.reader(brut.splitlines(), dialecte), 1) if any(r)]
        return "=== Feuille csv ===\n" + "\n".join(lignes), 1
    import pandas as pd

    sortie = []
    tables = pd.read_excel(chemin, sheet_name=None, header=None)
    for nom, table in tables.items():
        lignes = [f"L{n}: " + " | ".join(valeur(v) if v == v else "" for v in rangee)
                  for n, rangee in enumerate(table.itertuples(index=False), 1)]
        sortie.append(f"=== Feuille {nom} ===\n" + "\n".join(lignes))
    return "\n".join(sortie), len(tables)


def texte_courriel(chemin: Path) -> str:
    message = email.message_from_bytes(chemin.read_bytes(), policy=email.policy.default)
    entete = [f"{h}: {message[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if message[h]]
    corps = message.get_body(preferencelist=("plain", "html"))
    texte = corps.get_content() if corps else ""
    if corps is not None and corps.get_content_type() == "text/html":
        texte = html.unescape(re.sub(r"<[^>]+>", " ", texte))
    jointes = [p.get_filename() for p in message.iter_attachments() if p.get_filename()]
    return "\n".join(entete + ([f"Pièces jointes: {', '.join(jointes)}"] if jointes else [])) + "\n\n" + texte


def texte_traitement(chemin: Path) -> str:
    suffixe = chemin.suffix.lower()
    if shutil.which("pandoc") and suffixe in (".docx", ".odt", ".rtf"):
        return lancer(["pandoc", str(chemin), "-t", "plain", "--wrap=none"])
    if suffixe == ".docx":
        import docx
        d = docx.Document(str(chemin))
        parties = [p.text for p in d.paragraphs]
        parties += [" | ".join(cell.text for cell in rangee.cells) for table in d.tables for rangee in table.rows]
        return "\n".join(parties)
    if shutil.which("textutil"):
        return lancer(["textutil", "-convert", "txt", "-stdout", str(chemin)])
    raise ValueError("aucun convertisseur")


def extraire(chemin: Path) -> tuple[str, dict, str]:
    """(texte avec repères, {"pages"|"feuilles": n}, remarque)."""
    suffixe = chemin.suffix.lower()
    try:
        if suffixe == ".pdf":
            pages, remarque = pages_pdf(chemin)
            return avec_pages(pages), {"pages": len(pages)}, remarque
        if suffixe in TABLEURS:
            texte, n = feuilles_xlsx(chemin) if suffixe in (".xlsx", ".xlsm") else feuilles_autres(chemin)
            return texte, {"feuilles": n}, ""
        if suffixe in (".docx", ".doc", ".odt", ".rtf"):
            return texte_traitement(chemin), {}, ""
        if suffixe == ".pptx":
            import pptx
            diapositives = [[forme.text_frame.text for forme in d.shapes if forme.has_text_frame]
                            for d in pptx.Presentation(str(chemin)).slides]
            return avec_pages(["\n".join(d) for d in diapositives]), {"pages": len(diapositives)}, ""
        if suffixe == ".eml":
            return texte_courriel(chemin), {}, ""
        if suffixe in (".html", ".htm"):
            brut = chemin.read_text(encoding="utf-8", errors="replace")
            return html.unescape(re.sub(r"<[^>]+>", " ", brut)), {}, ""
        if suffixe in (".txt", ".md", ".json", ".xml"):
            return chemin.read_text(encoding="utf-8", errors="replace"), {}, ""
        if suffixe in (".png", ".jpg", ".jpeg", ".tif", ".tiff") and shutil.which("tesseract"):
            return avec_pages([ocr_image(chemin)]), {"pages": 1}, "texte obtenu par OCR : vérifier chaque extrait sur l'original"
        return "", {}, f"format {suffixe or 'inconnu'} non lu par le script (archive ? image ?) : l'ouvrir autrement"
    except Exception as exc:  # une pièce récalcitrante ne doit pas bloquer l'inventaire
        return "", {}, f"échec de l'extraction ({type(exc).__name__} : {str(exc)[:80]})"


def contenu(texte: str) -> list[str]:
    """Les mots du texte, sans les repères de page et de ligne."""
    lignes = [c.LIGNE_TABLEUR.sub("", ligne) for ligne in texte.splitlines() if not c.MARQUE.match(ligne.strip())]
    return " ".join(lignes).split()


def mots_du_nom(nom: str) -> str:
    return " ".join(re.sub(r"[_\-.()\[\]]+", " ", c.cle(Path(nom).stem)).split())


def cle_version(nom: str) -> str:
    sans_index = re.sub(r"^\d+[a-z]?\s+", "", mots_du_nom(nom))
    return " ".join(re.sub(rf"\b({VERSION})\b", " ", sans_index).split())


def jetons(nom: str) -> set[str]:
    return {m for m in mots_du_nom(nom).split()
            if len(m) >= 4 and not m.isdigit() and m not in GENERIQUES and not re.fullmatch(VERSION, m)}


def ngrammes(mots: list[str]) -> set[tuple]:
    return {tuple(mots[i:i + 5]) for i in range(len(mots) - 4)}


def differences(a: list[str], b: list[str]) -> str:
    ecarts = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op != "equal":
            ecarts.append(f"« {' '.join(a[i1:i2][:8]) or '—'} » → « {' '.join(b[j1:j2][:8]) or '—'} »")
    return " ; ".join(ecarts[:3]) + (f" (+{len(ecarts) - 3})" if len(ecarts) > 3 else "")


def signaler(pieces: list[dict], textes: dict[str, str]) -> None:
    par_id = {p["id"]: p for p in pieces}
    utiles = [p for p in pieces if p["mots"]]
    debut = {p["id"]: c.cle(" ".join(contenu(textes[p["id"]])[:60])) for p in utiles}

    def prefixer(piece, statut, raison):
        if not piece["statut"]:
            piece["statut"], piece["raison"] = statut, raison

    for p in pieces:
        if SYSTEME.match(p["nom"]):
            prefixer(p, "écartée", "fichier système")
        elif not p["mots"]:
            prefixer(p, "illisible", p.pop("_remarque", "") or "aucun texte")
        elif p.get("_remarque"):
            p["signaux"].append(p.pop("_remarque"))
        if p["format"] == ".pdf" and p["mots"] and p.get("pages") and p["mots"] / p["pages"] < 20:
            p["signaux"].append(f"peu de texte ({p['mots']} mots pour {p['pages']} pages) : scan partiel ? voir l'original")

    groupes = defaultdict(list)
    for p in pieces:
        groupes[p["sha256"]].append(p)
    for groupe in groupes.values():
        for p in groupe[1:]:
            p["signaux"].append(f"doublon exact de {groupe[0]['id']} ({groupe[0]['nom']})")
            prefixer(p, "écartée", f"doublon exact (même empreinte) de {groupe[0]['id']}")

    normes = defaultdict(list)
    for p in utiles:
        normes[c.normaliser(" ".join(contenu(textes[p["id"]])))].append(p)
    for groupe in normes.values():
        for p in groupe[1:]:
            if p["sha256"] != groupe[0]["sha256"]:
                p["signaux"].append(f"même texte que {groupe[0]['id']} dans un autre fichier : comparer format et signature")

    mots = {p["id"]: contenu(textes[p["id"]]) for p in utiles}
    ensembles = {i: ngrammes([m.lower() for m in w]) for i, w in mots.items() if len(w) >= 30}
    paires = set()
    for i, a in ensembles.items():
        proches = [(len(a & b) / len(a | b), j) for j, b in ensembles.items() if j != i and a and b]
        if not proches:
            continue
        score, j = max(proches)
        if (j, i) in paires:
            continue  # la paire est déjà signalée sur l'autre pièce
        paires.add((i, j))
        if 0.7 <= score < 1 and c.normaliser(" ".join(mots[i])) != c.normaliser(" ".join(mots[j])):
            par_id[i]["signaux"].append(f"texte proche de {j} ({round(score * 100)} %) — de {j} à {i} : {differences(mots[j], mots[i])}")

    versions = defaultdict(list)
    for p in pieces:
        if cle_version(p["nom"]):
            versions[cle_version(p["nom"])].append(p)
    for groupe in versions.values():
        for p in groupe:
            autres = [q["id"] for q in groupe if q is not p and q["sha256"] != p["sha256"]]
            if autres:
                p["signaux"].append(f"autre version du même document (d'après le nom) : {', '.join(autres)}")

    for p in utiles:
        nom, texte = mots_du_nom(p["nom"]), c.cle(" ".join(mots[p["id"]]))
        for source, trouve in (("le nom", PROJET_NOM.search(nom)), ("le début du texte", PROJET_DEBUT.search(debut[p["id"]])),
                               ("le texte", PROJET_TEXTE.search(texte))):
            if trouve:
                p["signaux"].append(f"projet ou version non signée ? « {trouve.group(0)} » dans {source}")
                break
        if p["format"] not in TABLEURS and ACTE.search(f"{nom} {debut[p['id']][:120]}") and not SIGNATURE.search(texte):
            p["signaux"].append("acte sans aucune mention de signature : vérifier qu'il est signé")
        if AVENANT.search(f"{nom} {debut[p['id']][:120]}"):
            p["signaux"].append(rattacher(p, pieces, textes))


def rattacher(avenant: dict, pieces: list[dict], textes: dict[str, str]) -> str:
    cherches = jetons(avenant["nom"])
    autres = [p for p in pieces if p is not avenant and not AVENANT.search(mots_du_nom(p["nom"]))]
    candidats = sorted(((len(cherches & jetons(p["nom"])), p["id"]) for p in autres), reverse=True)
    candidats = [i for score, i in candidats if score][:3]
    if candidats:
        return f"avenant ou annexe à rattacher à son contrat principal — candidats d'après le nom : {', '.join(candidats)}"
    cites = [p["id"] for p in autres if cherches and all(m in c.cle(textes.get(p["id"], "")) for m in cherches)][:4]
    return ("avenant ou annexe : contrat principal introuvable dans la data room d'après les noms — le demander"
            + (f" (« {' '.join(sorted(cherches))} » cité dans {', '.join(cites)})" if cites else ""))


def main() -> None:
    racine = Path(sys.argv[1]) if len(sys.argv) > 1 else c.DOCUMENTS
    if not racine.is_dir():
        raise SystemExit(f"{racine} n'est pas un dossier — indique le dossier des pièces")
    c.TEXTE.mkdir(parents=True, exist_ok=True)
    anciens = {p["piece"]: p for p in c.lire_jsonl(c.REGISTRE)[0]}
    fichiers = sorted((f for f in racine.rglob("*") if f.is_file()), key=lambda f: f.relative_to(racine).as_posix())
    pieces, textes = [], {}
    for n, chemin in enumerate(fichiers, 1):
        relatif = unicodedata.normalize("NFC", chemin.relative_to(racine).as_posix())
        texte, unites, remarque = ("", {}, "fichier système") if SYSTEME.match(chemin.name) else extraire(chemin)
        cible = c.TEXTE / (relatif + ".txt")
        cible.parent.mkdir(parents=True, exist_ok=True)
        cible.write_text(texte, encoding="utf-8")
        piece = {"id": f"P{n:03d}", "rubrique": relatif.split("/")[0] if "/" in relatif else "(racine)",
                 "piece": relatif, "source": str(chemin.absolute()), "nom": chemin.name,
                 "format": chemin.suffix.lower(), "taille": chemin.stat().st_size,
                 **unites, "mots": len(contenu(texte)), "sha256": hashlib.sha256(chemin.read_bytes()).hexdigest(),
                 "texte": cible.relative_to(c.ESPACE).as_posix(), "signaux": [], "statut": "", "raison": "",
                 "fiche": "", "affichee": False, "_remarque": remarque}
        ancien = anciens.get(relatif)
        if ancien and ancien.get("sha256") == piece["sha256"]:
            piece.update({k: ancien.get(k, piece[k]) for k in ("statut", "raison", "fiche", "affichee")})
        pieces.append(piece)
        textes[piece["id"]] = texte
    signaler(pieces, textes)
    for p in pieces:
        p.pop("_remarque", None)
    with c.verrou():
        c.ecrire_registre(pieces)
    rapport = resume(pieces, racine)
    (c.DD / "inventaire.md").write_text(rapport + "\n", encoding="utf-8")
    print(rapport)


def resume(pieces: list[dict], racine: Path) -> str:
    formats = defaultdict(int)
    rubriques = defaultdict(list)
    for p in pieces:
        formats[p["format"] or "sans extension"] += 1
        rubriques[p["rubrique"]].append(p)
    lignes = [f"{len(pieces)} pièces dans {racine}, {len(rubriques)} rubriques ("
              + ", ".join(f"{n} {f}" for f, n in sorted(formats.items(), key=lambda x: -x[1])) + ")",
              "Texte : dd/texte/   Registre : dd/registre.jsonl   Cet inventaire : dd/inventaire.md", ""]
    for rubrique, groupe in rubriques.items():
        lignes.append(f"{rubrique} — {len(groupe)} pièces")
        for p in groupe:
            volume = f"{p['pages']} p." if "pages" in p else f"{p['feuilles']} f." if "feuilles" in p else ""
            etat = f"  [{p['statut']}]" if p["statut"] else ""
            lignes.append(f"  {p['id']} {p['format'][1:]:4} {volume:>6} {p['mots']:>6} mots {-(-p['taille'] // 1024):>5} Ko  {p['nom'][:80]}{etat}")
    signales = [p for p in pieces if p["signaux"]]
    lignes += ["", f"SIGNAUX ({sum(len(p['signaux']) for p in signales)} sur {len(signales)} pièces) — chacun appelle une "
               "décision : écarter (raison), lire en le signalant, rattacher, ou demander la pièce manquante"]
    for p in signales:
        for signal in p["signaux"]:
            lignes.append(f"  {p['id']} {p['nom'][:50]} : {signal}")
    prefixes = defaultdict(int)
    for p in pieces:
        if p["statut"]:
            prefixes[p["statut"]] += 1
    lignes += ["", "Statuts déjà remplis : " + (", ".join(f"{n} {s}" for s, n in prefixes.items()) or "aucun")
               + f" ; {sum(1 for p in pieces if not p['statut'])} pièces à lire.",
               f"Suite : python {c.SCRIPTS}/lire.py <début du nom de rubrique, ex. 01>, rubrique par rubrique."]
    return "\n".join(lignes)


if __name__ == "__main__":
    main()
