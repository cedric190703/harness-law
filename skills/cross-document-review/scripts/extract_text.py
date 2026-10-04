#!/usr/bin/env python3
"""Convert every task document to plain text and print an inventory.

Usage: python extract_text.py [documents_dir] [out_dir]
Defaults: $DOCUMENTS_DIR and $WORKSPACE_DIR/text.
"""
import email
import email.policy
import html
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

from sections import find_sections


def run(cmd: list[str]) -> str:
    return subprocess.run(cmd, capture_output=True, text=True, check=True, timeout=120).stdout


def docx_text(path: Path) -> str:
    if shutil.which("pandoc"):
        return run(["pandoc", str(path), "-t", "plain", "--wrap=none"])
    import docx  # python-docx fallback

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def sheet_text(path: Path) -> str:
    import pandas as pd

    if path.suffix.lower() == ".csv":
        return pd.read_csv(path).to_string()
    out = []
    for name, df in pd.read_excel(path, sheet_name=None, header=None).items():
        out.append(f"### Sheet: {name}")
        out.append(df.fillna("").to_string(index=False, header=False))
    return "\n".join(out)


def pdf_text(path: Path) -> str:
    if shutil.which("pdftotext"):
        return run(["pdftotext", "-layout", str(path), "-"])
    import pdfplumber

    with pdfplumber.open(str(path)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def eml_text(path: Path) -> str:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    head = [f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h]]
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        text = html.unescape(re.sub(r"<[^>]+>", " ", text))
    attachments = [part.get_filename() for part in msg.iter_attachments() if part.get_filename()]
    if attachments:
        head.append("Attachments: " + ", ".join(attachments))
    return "\n".join(head) + "\n\n" + text


def pptx_text(path: Path) -> str:
    import pptx

    out = []
    for i, slide in enumerate(pptx.Presentation(str(path)).slides, 1):
        out.append(f"### Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n".join(out)


def markitdown_text(path: Path) -> str:
    from markitdown import MarkItDown

    return MarkItDown().convert(str(path)).text_content


CONVERTERS = {
    ".docx": docx_text,
    ".xlsx": sheet_text,
    ".xlsm": sheet_text,
    ".xls": sheet_text,
    ".csv": sheet_text,
    ".pdf": pdf_text,
    ".eml": eml_text,
    ".pptx": pptx_text,
}


def convert(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md", ".json", ".html", ".xml"):
        return path.read_text(encoding="utf-8", errors="replace")
    try:
        return CONVERTERS.get(suffix, markitdown_text)(path)
    except Exception as exc:  # keep going: one bad file must not block the inventory
        try:
            return markitdown_text(path)
        except Exception:
            return f"[extraction failed: {exc}]"


def main() -> None:
    docs = Path(sys.argv[1] if len(sys.argv) > 1 else os.environ.get("DOCUMENTS_DIR", "documents"))
    out = Path(sys.argv[2] if len(sys.argv) > 2 else Path(os.environ.get("WORKSPACE_DIR", ".")) / "text")
    out.mkdir(parents=True, exist_ok=True)
    files = sorted(p for p in docs.rglob("*") if p.is_file())
    print(f"{len(files)} documents in {docs}\n")
    print(f"{'document':60} {'words':>7} {'sections':>8}")
    for path in files:
        text = convert(path)
        rel = path.relative_to(docs)
        target = out / (str(rel) + ".txt")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8")
        print(f"{str(rel):60} {len(text.split()):>7} {len(find_sections(text)):>8}")
    print(f"\nText written to {out}/<document>.txt — read every document in full before building the ledger.")


if __name__ == "__main__":
    main()
