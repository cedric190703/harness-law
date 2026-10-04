#!/usr/bin/env python3
"""Refaire les totaux des tableurs de la data room et signaler les écarts.

  python tableaux.py              # tous les tableurs du registre
  python tableaux.py P045 P046    # ceux-là

Chaque ligne « Total / Sous-total / Somme » est comparée à la somme des lignes au-dessus, et
chaque colonne intitulée « Total » à la somme des cellules à sa gauche. Signale aussi les totaux
saisis en dur dans une feuille qui contient des formules. Un écart est une piste, pas une
conclusion : relis la feuille (lignes « L… » du texte) et refais le calcul avant d'en faire un constat.
"""
import re
import sys
from pathlib import Path

import commun as c

TOTAL = re.compile(r"^(sous ?-? ?)?total|^somme|^cumul")
TOLERANCE = 0.01


def nombre(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def montant(v: float) -> str:
    return f"{v:,.2f}".replace(",", " ").replace(".", ",")


def egal(a: float, b: float) -> bool:
    return abs(a - b) <= TOLERANCE + 1e-9 * abs(b)


def est_total(v) -> bool:
    return isinstance(v, str) and bool(TOTAL.match(c.cle(v)))


def sommes_partielles(valeurs: list[float]) -> list[float]:
    """Sommes de 2, 3, … valeurs consécutives en partant de la plus proche du total."""
    sommes, cumul = [], 0.0
    for i, v in enumerate(valeurs):
        cumul += v
        if i:
            sommes.append(cumul)
    return sommes


def colonne(j: int) -> str:
    lettres, j = "", j + 1
    while j:
        j, reste = divmod(j - 1, 26)
        lettres = chr(65 + reste) + lettres
    return lettres


def ecart(lieu: str, declare: float, attendu: float, comment: str) -> str:
    part = f" ({(declare - attendu) / attendu * 100:+.1f} %)".replace(".", ",") if attendu else ""
    signe = "+" if declare > attendu else ""
    return f"{lieu} : déclaré {montant(declare)} ; {comment} = {montant(attendu)} ; écart {signe}{montant(declare - attendu)}{part}"


def controler_feuille(valeurs: list[list], formules: list[list]) -> tuple[list[str], list[str]]:
    """(écarts, remarques) d'une feuille."""
    ecarts, remarques = [], []
    cellule = lambda grille, r, j: grille[r][j] if j < len(grille[r]) else None  # noqa: E731
    sans_valeur = lambda r, j: cellule(valeurs, r, j) is None and str(cellule(formules, r, j) or "").startswith("=")  # noqa: E731
    a_des_formules = any(str(f).startswith("=") for rangee in formules for f in rangee if f is not None)
    lignes_total = {r for r, rangee in enumerate(valeurs)
                    if any(est_total(v) for v in rangee[:3] if isinstance(v, str) and v.strip())}
    for r in sorted(lignes_total):
        en_dur = []
        for j, v in enumerate(valeurs[r]):
            declare = nombre(v)
            if declare is None:
                continue
            if a_des_formules and not str(cellule(formules, r, j)).startswith("="):
                en_dur.append(colonne(j))
            details, totaux, plage, incomplet = [], [], [], False
            for k in range(r - 1, -1, -1):
                v_k = cellule(valeurs, k, j)
                if isinstance(v_k, str) and v_k.strip():
                    break  # en-tête de colonne : début du tableau
                incomplet |= sans_valeur(k, j)
                if nombre(v_k) is None:
                    continue
                (totaux if k in lignes_total else details).append(nombre(v_k))
                plage += [] if k in lignes_total else [k]
            if incomplet:
                remarques.append(f"L{r + 1} colonne {colonne(j)} : formules sans valeur calculée au-dessus — refais le total à la main")
            elif len(details) + len(totaux) >= 2 and not any(egal(declare, s) for s in sommes_partielles(details) + [sum(totaux)] * bool(totaux)):
                lignes = f"L{min(plage) + 1}–L{max(plage) + 1}" if plage else "des lignes au-dessus"
                ecarts.append(ecart(f"L{r + 1} colonne {colonne(j)}", declare, sum(details), f"somme {lignes}"))
        if en_dur:
            remarques.append(f"L{r + 1} : total saisi en dur (colonnes {', '.join(en_dur)}) alors que la feuille contient des formules")
    for h, entete in enumerate(valeurs):
        colonnes_total = [j for j, v in enumerate(entete) if est_total(v)]
        if not colonnes_total or h in lignes_total or sum(isinstance(v, str) for v in entete) < 2:
            continue
        for r in range(h + 1, len(valeurs)):
            if r not in lignes_total and any(est_total(v) for v in valeurs[r]):
                break  # un nouveau tableau commence
            for j in colonnes_total:
                declare = nombre(cellule(valeurs, r, j))
                if declare is None or any(sans_valeur(r, x) for x in range(j)):
                    continue
                suite = []
                for v in reversed(valeurs[r][:j]):
                    if nombre(v) is None:
                        break
                    suite.append(nombre(v))
                if len(suite) >= 2 and not any(egal(declare, s) for s in sommes_partielles(suite)):
                    ecarts.append(ecart(f"L{r + 1} colonne {colonne(j)} (« {entete[j]} »)", declare, sum(suite),
                                        f"somme des {len(suite)} cellules à gauche"))
    return ecarts, remarques


def main() -> None:
    import openpyxl

    pieces = c.lire_registre()
    choisies = [c.resoudre(pieces, a) for a in sys.argv[1:]] or [p for p in pieces if p["format"] in (".xlsx", ".xlsm")]
    total, sans = 0, []
    for p in (p for p in choisies if p):
        try:
            chemin = Path(p.get("source") or c.DOCUMENTS / p["piece"])
            classeur_v = openpyxl.load_workbook(chemin, data_only=True)
            classeur_f = openpyxl.load_workbook(chemin)
        except Exception as exc:
            print(f"\n{p['id']} {p['nom']} : illisible par openpyxl ({type(exc).__name__}) — refais les calculs à la main")
            continue
        lignes = []
        for feuille in classeur_v.worksheets:
            valeurs = [[x.value for x in rangee] for rangee in feuille.iter_rows()]
            formules = [[x.value for x in rangee] for rangee in classeur_f[feuille.title].iter_rows()]
            ecarts, remarques = controler_feuille(valeurs, formules)
            total += len(ecarts)
            lignes += [f"  feuille « {feuille.title} » {e}" for e in ecarts]
            lignes += [f"  feuille « {feuille.title} » (remarque) {r}" for r in remarques]
        if lignes:
            print(f"\n{p['id']} {p['piece']}")
            print("\n".join(lignes))
        else:
            sans.append(p["id"])
    print(f"\n{total} écart(s) de total à vérifier ; {len(sans)} tableur(s) sans écart ni remarque"
          + (f" ({', '.join(sans)})" if sans else "") + ".")
    print("Ce script ne voit que les totaux. Ratios, covenants, chiffres repris d'une pièce à l'autre : "
          "refais-les avec python (openpyxl ou pandas) et note le calcul dans le champ « calcul » du constat.")


if __name__ == "__main__":
    main()
